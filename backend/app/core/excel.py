"""Shared Excel import helpers: RTL templates, row reading and Persian-friendly parsing.

Two steps everywhere: preview (row-by-row validation, nothing saved) and commit
(all-or-nothing; refused while any row has an error).
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel

from app.core.errors import DomainError
from app.core.money import _FA, MONEY_PLACES

MAX_ROWS = 50_000
YES = {"بله", "آری", "yes", "y", "1", "true", "دارد", "✓"}
NO = {"خیر", "نه", "no", "n", "0", "false", "ندارد", ""}


class RowError(BaseModel):
    row: int
    column: str
    message: str


class ImportPreview(BaseModel):
    total_rows: int
    valid_rows: int
    to_create: int
    to_update: int
    errors: list[RowError]
    sample: list[dict[str, Any]]


class ImportResult(BaseModel):
    created: int
    updated: int
    document_id: uuid.UUID | None = None


def template(
    columns: list[str],
    title: str,
    example: list[Any],
    note: str,
    *,
    text_columns: tuple[str, ...] = ("A",),
    yes_no_columns: tuple[str, ...] = (),
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.append(columns)
    ws.append(example)
    head = PatternFill("solid", fgColor="5146E5")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, _ in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 22
    # Codes, barcodes and IDs as text, so Excel keeps leading zeros and long digit strings.
    for col in text_columns:
        ws.column_dimensions[col].number_format = "@"
        ws[f"{col}2"].number_format = "@"
    if yes_no_columns:
        dv = DataValidation(type="list", formula1='"بله,خیر"', allow_blank=True)
        ws.add_data_validation(dv)
        for col in yes_no_columns:
            dv.add(f"{col}2:{col}{MAX_ROWS + 1}")
    info = wb.create_sheet("راهنما")
    info.sheet_view.rightToLeft = True
    for line in note.splitlines():
        info.append([line])
    info.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_rows(file_bytes: bytes, ncols: int) -> list[tuple[int, list[Any]]]:
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - any parse failure is a user error here
        raise DomainError("فایل Excel قابل خواندن نیست (فقط ‎.xlsx)", code="invalid_file") from e
    ws = wb.worksheets[0]
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, max_col=ncols, values_only=True), start=2):
        values = list(row) + [None] * (ncols - len(row))
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        rows.append((idx, values))
        if len(rows) > MAX_ROWS:
            raise DomainError(f"حداکثر {MAX_ROWS} ردیف در هر فایل", code="too_many_rows")
    wb.close()
    return rows


def text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip().translate(_FA)


def decimal(v: Any, places: str) -> Decimal:
    s = text(v).replace(",", "").replace("٬", "").replace("٫", ".")
    try:
        d = Decimal(s)
    except InvalidOperation as e:
        raise ValueError("عدد نامعتبر") from e
    q = Decimal(places)
    if d != d.quantize(q):
        raise ValueError(f"حداکثر {len(places.split('.')[1]) if '.' in places else 0} رقم اعشار")
    return d.quantize(q)


def to_rials(v: Any, currency: str) -> Decimal:
    d = decimal(v, "0.0001" if currency == "IRR" else "0.001")
    if d < 0:
        raise ValueError("منفی مجاز نیست")
    return (d * 10 if currency == "IRT" else d).quantize(MONEY_PLACES)


def yes_no(v: Any) -> bool:
    t = text(v).lower()
    if t in YES:
        return True
    if t in NO:
        return False
    raise ValueError("«بله» یا «خیر»")


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD = 25 * 1024 * 1024


def upload_bytes() -> bytes:
    """The uploaded .xlsx of the current request (multipart field ``file``)."""
    from flask import request

    request.max_content_length = MAX_UPLOAD
    f = request.files.get("file")
    if f is None:
        raise DomainError("فایل Excel را انتخاب کنید", code="file_required")
    data: bytes = f.read()
    if not data.startswith(b"PK"):
        raise DomainError("فقط فایل ‎.xlsx پذیرفته می‌شود", code="invalid_file")
    return data
