"""Excel export of any ReportTable (RTL sheet, amounts in the display currency).

PDF is produced by the browser's print view ("Save as PDF"): server-side PDF with
correctly shaped Persian text would need another rendering stack for no gain.
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.modules.reports.schemas import ReportTable


def _cell(kind: str, value: Any, currency: str) -> Any:
    if value is None or value == "":
        return None
    if kind == "money":
        d = Decimal(value)
        return d / 10 if currency == "IRT" else d
    if kind in ("qty", "percent"):
        return Decimal(value)
    return value


def xlsx(table: ReportTable, currency: str, filters: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = table.title[:31]
    ws.sheet_view.rightToLeft = True
    unit = "تومان" if currency == "IRT" else "ریال"
    ws.append([table.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([filters])
    ws.append([])
    header = [f"{c.label} ({unit})" if c.kind == "money" else c.label for c in table.columns]
    ws.append(header)
    fill = PatternFill("solid", fgColor="5146E5")
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    money_fmt = "#,##0.#" if currency == "IRT" else "#,##0"
    for row in table.rows:
        values = [_cell(c.kind, row.get(c.key), currency) for c in table.columns]
        if row.get("_level"):
            values[1 if len(values) > 1 else 0] = "    " * int(row["_level"]) + str(values[1] or "")
        ws.append(values)
        if row.get("_bold"):
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    if table.totals:
        ws.append([_cell(c.kind, table.totals.get(c.key), currency) for c in table.columns])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for i, col in enumerate(table.columns, start=1):
        letter = ws.cell(4, i).column_letter
        ws.column_dimensions[letter].width = 34 if col.kind == "text" else 18
        if col.kind == "money":
            for cell in ws[letter][4:]:
                cell.number_format = money_fmt
    for n in table.notes:
        ws.append([])
        ws.append([n])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
